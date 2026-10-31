import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from vatty import apresentacao as a
from vatty import cli

ROTEIRO = """# Vatty: ouve e transforma
subtitulo: Tudo no seu PC.

## Como funciona
tipo: passos
- Grava
- Transcreve
horario: [00:55]

## O painel
tipo: imagem
- Gravando
print: frames/0001.jpg
recorte: 20 10 60 80
notas: mostrar o REC

## Estranho
tipo: bolo
recorte: 90 0 20 50
- <script>alert(1)</script> e **forte**
"""


class TestLer(unittest.TestCase):
    def test_roteiro_vira_capa_e_slides(self):
        deck = a.ler(ROTEIRO)
        self.assertEqual((deck["titulo"], deck["subtitulo"]), ("Vatty: ouve e transforma", "Tudo no seu PC."))
        self.assertEqual([s["tipo"] for s in deck["slides"]], ["passos", "imagem", "lista"])  # tipo desconhecido vira lista
        self.assertEqual(deck["slides"][1]["recorte"], (20, 10, 60, 80))
        self.assertEqual(deck["slides"][1]["notas"], "mostrar o REC")

    def test_recorte_fora_da_imagem_e_ignorado(self):
        self.assertIsNone(a.ler(ROTEIRO)["slides"][2]["recorte"])

    def test_sem_titulo_ou_sem_slide_da_erro(self):
        with self.assertRaises(a.ErroApresentacao):
            a.ler("## Só um slide\n- item")
        with self.assertRaises(a.ErroApresentacao):
            a.ler("# Só o título")


class TestHtml(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.pasta = Path(tmp.name)
        (self.pasta / "frames").mkdir()
        (self.pasta / "frames" / "0001.jpg").write_bytes(b"\xff\xd8jpg")

    def test_escapa_texto_embute_o_print_e_conta_os_slides(self):
        doc = a.para_html(a.ler(ROTEIRO), self.pasta, arte=None)
        self.assertNotIn("<script>alert", doc)
        self.assertIn("&lt;script&gt;", doc)
        self.assertIn("<strong>forte</strong>", doc)
        self.assertIn("data:image/jpeg;base64,", doc)
        self.assertEqual(doc.count('<section class="slide'), 4)   # capa + 3

    def test_print_fora_da_pasta_nao_entra(self):
        deck = a.ler(ROTEIRO.replace("frames/0001.jpg", "../../etc/passwd"))
        self.assertNotIn("data:image", a.para_html(deck, self.pasta, arte=None))

    def test_youtube_vira_link_no_horario(self):
        fonte = {"tipo": "youtube", "url": "https://www.youtube.com/watch?v=abc"}
        self.assertIn("watch?v=abc&amp;t=55s", a.para_html(a.ler(ROTEIRO), self.pasta, fonte, arte=None))


class TestMontar(unittest.TestCase):
    def test_cada_roteiro_vira_html_ao_lado(self):
        with tempfile.TemporaryDirectory() as d:
            pasta = Path(d)
            (pasta / "fonte.json").write_text(json.dumps({"tipo": "gravacao", "data": "2026-10-31_13-43"}))
            (pasta / "apresentacao-pitch.md").write_text(ROTEIRO, encoding="utf-8")
            feitos = a.montar(pasta, formatos=("html",))
            self.assertEqual(feitos, [pasta / "apresentacao-pitch.html"])
            self.assertIn("31/10/2026", feitos[0].read_text(encoding="utf-8"))

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx não instalado")
    def test_pptx_tem_capa_slides_e_notas(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            pasta = Path(d)
            (pasta / "apresentacao-pitch.md").write_text(ROTEIRO, encoding="utf-8")
            feitos = a.montar(pasta)   # padrão: .pptx
            self.assertEqual(feitos, [pasta / "apresentacao-pitch.pptx"])
            prs = Presentation(feitos[0])
            self.assertEqual(len(prs.slides), 4)
            self.assertEqual(prs.slides[2].notes_slide.notes_text_frame.text, "mostrar o REC")
            textos = " ".join(f.text_frame.text for f in prs.slides[1].shapes if f.has_text_frame)
            self.assertIn("Como funciona", textos)

    def test_cli_sem_roteiro_da_erro(self):
        with tempfile.TemporaryDirectory() as d:
            self.assertEqual(cli.main(["apresentacao", d]), 1)


if __name__ == "__main__":
    unittest.main()
